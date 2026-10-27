"""Unit: budget suggestions applied by reply. Figures are parsed and bounded in Python, only one
input cell changes, a dated backup comes first, and anything unclear changes nothing. Runs on a
temp copy of the example workbook; the committed example is never touched.
Run with: PYTHONPATH=src .venv/bin/python -m unittest discover tests"""

import json
import shutil
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path
from unittest import mock

import openpyxl
import yaml

from butler import budget, budget_changes
from butler.budget_changes import parse_amounts, within_bounds

REPO = Path(__file__).resolve().parents[1]
TODAY = date(2026, 10, 2)
# Invented: the look-back on 1 October suggested these for the Johnsons.
OPEN = {"Transport": {"budget": 120.0, "trend": 185.0, "suggested": 175},
        "Food": {"budget": 950.0, "trend": 1100.0, "suggested": 1100}}
SETTINGS = {"email": {"trusted_senders": ["sam@example.com"]}}


def _action(category, amount=None, text="apply it"):
    return {"type": "budget_change", "text": text, "category": category, "amount": amount}


class ParseTest(unittest.TestCase):
    def test_reads_plain_and_grouped_figures(self):
        cases = {"160": 160, "€160": 160, "Transport 160.": 160, "1,050": 1050, "1.050": 1050,
                 "1,050.50": 1050.5, "1.050,50": 1050.5, "87,5": 87.5, "87.50": 87.5}
        for text, want in cases.items():
            self.assertEqual(parse_amounts(text), [want], text)
        self.assertEqual(parse_amounts("1 050"), [1, 50], "a plain space separates two figures")
        nbsp = {"1\u00a0050": 1050, "1\u202f050,50": 1050.5, "160\u00a0€": 160, "12\u00a0345\u00a0678": 12345678}
        for text, want in nbsp.items():
            self.assertEqual(parse_amounts(text), [want], repr(text))

    def test_unclear_figures_are_none(self):
        for text in ("1,05,0", "1,050.000", "12.345.67", "1\u00a01\u00a0000", "1234\u00a0567", "1\u00a05"):
            self.assertEqual(parse_amounts(text), [None], text)

    def test_every_figure_is_found(self):
        self.assertEqual(parse_amounts("Food: Groceries 800, not 900"), [800, 900])
        self.assertEqual(parse_amounts("yes please"), [])


class BoundsTest(unittest.TestCase):
    def test_between_half_the_lower_and_double_the_higher(self):
        self.assertTrue(within_bounds(175, 120, 185))
        self.assertTrue(within_bounds(60, 120, 185))
        self.assertTrue(within_bounds(370, 120, 185))
        self.assertFalse(within_bounds(59, 120, 185))
        self.assertFalse(within_bounds(371, 120, 185))
        self.assertFalse(within_bounds(0, 0, 10))
        self.assertFalse(within_bounds(-5, 120, 185))


class WorkbookCase(unittest.TestCase):
    """A throwaway install: a copy of the example workbook as budget/budget.xlsx, the example
    category map, euros, and the invented open suggestions above."""

    def setUp(self):
        self.root = Path(tempfile.mkdtemp(prefix="butler-budget-"))
        (self.root / "budget").mkdir()
        (self.root / "data").mkdir()
        self.book = self.root / "budget" / "budget.xlsx"
        shutil.copy(REPO / "budget" / "budget.example.xlsx", self.book)
        self.example_bytes = (REPO / "budget" / "budget.example.xlsx").read_bytes()
        rules = yaml.safe_load((REPO / "config" / "category_map.example.yaml").read_text())
        self.suggestions = self.root / "data" / "budget_suggestions.json"
        self._open(OPEN, TODAY - timedelta(days=1))
        for patch in (mock.patch.object(budget, "BUDGET_FILE", self.book),
                      mock.patch.object(budget, "ROOT", self.root),
                      mock.patch.object(budget_changes, "BUDGET_FILE", self.book),
                      mock.patch.object(budget_changes, "ROOT", self.root),
                      mock.patch.object(budget_changes, "SUGGESTIONS", self.suggestions),
                      mock.patch("butler.finance.load_rules", return_value=rules),
                      mock.patch("butler.fmt._symbols", return_value=("€", "EUR", "USD"))):
            patch.start()
            self.addCleanup(patch.stop)
        self.addCleanup(shutil.rmtree, self.root, True)

    def tearDown(self):
        self.assertEqual((REPO / "budget" / "budget.example.xlsx").read_bytes(), self.example_bytes,
                         "a test changed the committed example workbook")

    def _open(self, suggestions, sent):
        self.suggestions.write_text(json.dumps({"month": "2026-09", "sent": sent.isoformat(),
                                                "suggestions": suggestions}))

    def amount(self, item):
        ws = openpyxl.load_workbook(self.book)["Expenses"]
        return next(ws[f"D{r}"].value for r in range(5, 61) if ws[f"B{r}"].value == item)

    def history(self):
        folder = self.root / "budget" / "history"
        return sorted(p.name for p in folder.iterdir()) if folder.exists() else []


class PlanTest(WorkbookCase):
    def test_taking_a_suggestion_as_is(self):
        c = budget_changes.plan(_action("Transport"), "Yes, apply the Transport suggestion", TODAY)
        self.assertEqual((c.item, c.old, c.new, c.total_new), ("Transit and taxis", 120, 175, 175))

    def test_own_figure(self):
        c = budget_changes.plan(_action("Transport", "160"), "Transport 160 please", TODAY)
        self.assertEqual(c.new, 160)

    def test_figure_must_be_in_the_reply(self):
        said = budget_changes.plan(_action("Transport", "200"), "Transport, yes", TODAY)
        self.assertIsInstance(said, str)
        self.assertIn("couldn't find", said)

    def test_other_currency_or_several_figures_ask(self):
        for amount in ("$160", "160 USD", "150 or 160"):
            self.assertIsInstance(budget_changes.plan(_action("Transport", amount), f"Transport {amount}", TODAY),
                                  str, amount)

    def test_out_of_bounds_asks(self):
        said = budget_changes.plan(_action("Transport", "5000"), "Transport 5000", TODAY)
        self.assertIn("outside what I'll change by reply", said)
        said = budget_changes.plan(_action("Transport", "10"), "Transport 10", TODAY)
        self.assertIn("outside", said)

    def test_a_category_with_several_lines_needs_one_named(self):
        said = budget_changes.plan(_action("Food"), "Apply the Food one", TODAY)
        self.assertIn("Groceries, Restaurants", said)
        c = budget_changes.plan(_action("Food"), "Food: put it on Groceries", TODAY)
        self.assertEqual((c.item, c.old, c.new, c.total_old, c.total_new), ("Groceries", 700, 850, 950, 1100))
        c = budget_changes.plan(_action("Food", "800"), "Food: Groceries 800", TODAY)
        self.assertEqual((c.item, c.new, c.total_new), ("Groceries", 800, 1050))

    def test_which_suggestion_must_be_clear(self):
        self.assertIn("Open ones", budget_changes.plan(_action("Housing"), "Housing 2000", TODAY))
        self.assertIn("Open ones", budget_changes.plan(_action("Transport"), "Yes, do it", TODAY))
        self._open({"Transport": OPEN["Transport"]}, TODAY)
        self.assertEqual(budget_changes.plan(_action("Transport"), "Yes, do it", TODAY).new, 175)

    def test_old_suggestions_are_closed(self):
        self._open(OPEN, TODAY - timedelta(days=budget_changes.OPEN_DAYS + 1))
        self.assertIn("no open budget suggestions", budget_changes.plan(_action("Transport"), "Transport", TODAY))

    def test_second_currency_line_converts_in_whole_units(self):
        wb = openpyxl.load_workbook(self.book)
        wb["Expenses"]["C12"] = "USD"      # Transit and taxis, priced in dollars for this test
        wb.save(self.book)
        c = budget_changes.plan(_action("Transport"), "Transport", TODAY)
        self.assertEqual(c.new, round((175 - 0) / 0.86))
        self.assertAlmostEqual(c.total_new, c.new * 0.86)


class ApplyTest(WorkbookCase):
    def test_backup_first_then_one_cell_changes(self):
        before = self.book.read_bytes()
        said = budget_changes.apply(SETTINGS, _action("Transport"), "apply the Transport suggestion", TODAY)
        self.assertEqual(said, "Budget changed: Transit and taxis (Transport) from €120 to €175 a month. "
                               "The workbook as it was is saved as "
                               f"{Path('budget', 'history', 'budget_2026-10-02_reply.xlsx')}.")
        self.assertEqual(self.amount("Transit and taxis"), 175)
        self.assertEqual((self.root / "budget" / "history" / "budget_2026-10-02_reply.xlsx").read_bytes(), before)
        old = openpyxl.load_workbook(self.root / "budget" / "history" / "budget_2026-10-02_reply.xlsx")
        new = openpyxl.load_workbook(self.book)
        changed = [(ws.title, c.coordinate) for ws in new for row in ws.iter_rows() for c in row
                   if c.value != old[ws.title][c.coordinate].value]
        self.assertEqual(changed, [("Expenses", "D12")])

    def test_a_second_change_the_same_day_keeps_both_backups(self):
        budget_changes.apply(SETTINGS, _action("Transport"), "Transport", TODAY)
        budget_changes.apply(SETTINGS, _action("Transport", "150"), "Transport 150", TODAY)
        self.assertEqual(self.history(), ["budget_2026-10-02_reply.xlsx", "budget_2026-10-02_reply_2.xlsx"])
        self.assertEqual(self.amount("Transit and taxis"), 150)

    def test_several_lines_say_what_the_category_becomes(self):
        said = budget_changes.apply(SETTINGS, _action("Food"), "Food on Groceries", TODAY)
        self.assertIn("Groceries (Food) from €700 to €850 a month, so Food goes from €950 to €1,100", said)

    def test_anything_unclear_changes_nothing(self):
        before = self.book.read_bytes()
        for action, text in ((_action("Food"), "Food"), (_action("Transport", "5000"), "Transport 5000"),
                             (_action("Transport", "200"), "Transport"), (_action("Rent"), "Rent")):
            said = budget_changes.apply(SETTINGS, action, text, TODAY)
            self.assertTrue(said.startswith("I didn't change the budget."), said)
        self.assertEqual(self.book.read_bytes(), before)
        self.assertEqual(self.history(), [])

    def test_never_edits_the_example(self):
        self.book.unlink()
        said = budget_changes.apply(SETTINGS, _action("Transport"), "Transport", TODAY)
        self.assertIn("there isn't one yet", said)
        self.assertEqual(self.history(), [])

    def test_refuses_a_line_edited_since_it_was_read(self):
        with self.assertRaises(ValueError):
            budget.set_monthly_amount("Transport", "Transit and taxis", 999, 175, TODAY)
        self.assertEqual(self.amount("Transit and taxis"), 120)


class LookBackWordingTest(WorkbookCase):
    """The look-back only offers "reply to apply" when a reply can actually apply it."""

    def _sections(self, by_reply):
        from types import SimpleNamespace

        from butler import monthly
        mo = SimpleNamespace(
            start=date(2026, 9, 1), end=date(2026, 9, 30), done_steps=[], done_reminders=[], shopping_done=0,
            time={}, networth_change=None, runway_months=None, savings_balance=None,
            money=SimpleNamespace(one_offs_mtd=[], draw_mtd=0, rental_mtd=0, ledgers=[], managed_fees_mtd=0),
            suggestions=[monthly.Suggestion("Transport", 120, 186, {"2026-08": 180, "2026-09": 190}),
                         monthly.Suggestion("Food", 950, 1090, {"2026-08": 1050, "2026-09": 1110})])
        words = {"headline": "", "accomplishments": "", "time_insight": "", "money_insight": "", "next_month": []}
        with mock.patch("butler.render.split_lines", return_value=[]):
            return dict(monthly.sections(mo, words, by_reply=by_reply))

    def test_offers_reply_when_usable(self):
        secs = self._sections(True)
        lines = secs["Budget suggestions (reply to apply one)"]
        self.assertIn("suggest €175", lines[0])
        self.assertIn("Food has 2 lines (Groceries, Restaurants): name the one to change", lines[1])
        self.assertEqual(lines[-1], "To apply one, reply with its category, like “Transport”, or give your "
                                    "own monthly figure, like “Transport 175”. I save a copy of the budget "
                                    "workbook before changing it.")

    def test_says_edit_the_workbook_otherwise(self):
        secs = self._sections(False)
        lines = secs["Budget suggestions (edit the budget workbook to apply any)"]
        self.assertEqual(len(lines), 2)
        self.assertNotIn("reply", " ".join(lines))

    def test_usable_needs_the_households_own_workbook_and_trusted_senders(self):
        self.assertTrue(budget_changes.usable(SETTINGS))
        self.assertFalse(budget_changes.usable({"email": {"trusted_senders": []}}))
        self.book.unlink()
        self.assertFalse(budget_changes.usable(SETTINGS))


if __name__ == "__main__":
    unittest.main()
