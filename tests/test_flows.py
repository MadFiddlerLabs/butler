"""Integration: every email flow and every portal page, end to end, with Gmail, Claude, Lunch
Money, the news, the weather and the Calendar helper faked (tests/fakes). No network.
Run with: PYTHONPATH=src .venv/bin/python -m unittest discover tests"""

import json
import re
import unittest
from datetime import timedelta
from pathlib import Path

from butler.agenda import OUTSIDE_LABEL, OUTSIDE_MARK
from butler.fmt import strftime
from fakes import fixtures
from fakes.home import Home

TRUSTED_SUBJECT = "Re: Bento · Saturday"
ACTIONS = {TRUSTED_SUBJECT: [
    {"type": "add_shopping", "text": "milk"},
    {"type": "add_todo", "text": "Call the plumber"},
    {"type": "add_event", "text": "Dentist for Mia", "date": "2026-10-15", "start_time": "10:00",
     "end_time": None, "location": "Clínica Exemplo"},
    {"type": "step_done", "text": "Boiler serviced", "step_id": "boiler-service"},
    {"type": "household_note", "text": "Pat is away next Friday"},
    {"type": "feedback", "text": "Shorter news please"},
]}
SAFE_EK_COMMANDS = {"calendars", "events", "reminder-lists", "reminders", "ensure-list", "add-reminder",
                    "complete-reminder", "completed", "add-event"}


def _mentions(value, title: str, labelled: bool = False):
    """For each string in a data block that holds `title`: whether it's labelled as an invite from
    outside, by its own text or by a "from" beside it."""
    if isinstance(value, dict):
        labelled = labelled or value.get("from") == OUTSIDE_LABEL
        for v in value.values():
            yield from _mentions(v, title, labelled)
    elif isinstance(value, list):
        for v in value:
            yield from _mentions(v, title, labelled)
    elif isinstance(value, str) and title in value:
        yield labelled or OUTSIDE_MARK in value or OUTSIDE_LABEL in value


class FullHouseholdTest(unittest.TestCase):
    """One household with everything connected, run through each flow in order."""

    @classmethod
    def setUpClass(cls):
        cls.home = Home()
        h = cls.home
        cls.trusted_id = h.add_reply("sam@example.com", TRUSTED_SUBJECT,
                                     "add milk to shopping\nand call the plumber", "example.com")
        h.add_reply("stranger@attacker.test", "Re: Bento", "add a pony to shopping", "attacker.test")
        h.add_reply("alex@example.com", "Re: Bento", "add caviar", "example.com.attacker.test")
        h.add_reply("alex@example.com", "Re: Bento", "add champagne", None)
        h.set_actions(ACTIONS)
        cls.daily = h.run("daily")
        cls.replies = h.run("replies")
        cls.replies_again = h.run("replies")
        cls.weekly = h.run("weekly")
        cls.monthly = h.run("monthly")
        cls.portal = h.run("portal")

    @classmethod
    def tearDownClass(cls):
        cls.home.close()

    def test_no_flow_crashed_or_broke_a_schema(self):
        for name in ("daily", "replies", "weekly", "monthly", "portal"):
            record = getattr(self, name)
            self.assertIsNone(record["error"], f"{name}:\n{record['error']}")
            self.assertEqual(record["schema_errors"], [], name)

    # ---------- the morning brief ----------

    def test_brief_is_sent_once_to_the_household(self):
        mails = self.daily["mails"]
        self.assertEqual(len(mails), 1)
        m = mails[0]
        self.assertTrue(m["subject"].startswith("Bento · "), m["subject"])
        self.assertIn("Bento", m["from"])
        self.assertEqual(m["to"], ["sam@example.com", "alex@example.com"])
        self.assertTrue(m["text"].strip() and m["html"].strip())
        self.assertEqual(self.daily["smtp_logins"], [{"user": "bento.johnson.butler@example.com", "password_ok": True}])

    def test_brief_has_every_section_and_no_problems(self):
        m = self.daily["mails"][0]
        for expected in ("Mia school drop-off", "Renew the parking permit", "Invented headline",
                         "Lisbon today", "Bring the renewal form"):
            self.assertIn(expected, m["text"], expected)
        self.assertNotIn("Problems this morning", m["text"])
        self.assertNotIn("Not connected yet", m["text"])
        self.assertTrue(m["images"], "the brief has no charts")

    def test_an_invite_from_outside_is_marked_quietly(self):
        m = self.daily["mails"][0]
        self.assertIn("Dinner with the Costas (invite from outside)", m["text"])
        self.assertIn("Dinner with the Costas<span", m["html"])
        for own in ("standup", "school drop-off", "swimming"):
            self.assertNotIn(f"{own} (invite from outside)", m["text"])

    def test_outside_titles_reach_claude_only_labelled_inside_a_data_block(self):
        title = "Dinner with the Costas"
        prompts = [p for flow in ("daily", "weekly", "monthly") for p in getattr(self, flow)["prompts"]]
        self.assertTrue(any(title in p["prompt"] for p in prompts))
        for p in prompts:
            self.assertNotIn(title, p["system"], p["schema"])
            outside_blocks = re.sub(r"<<<\w+\n.*?\n>>>", "", p["prompt"], flags=re.S)
            self.assertNotIn(title, outside_blocks, p["schema"])
            for found in re.findall(r"<<<\w+\n(.*?)\n>>>", p["prompt"], flags=re.S):
                for labelled in _mentions(json.loads(found), title):
                    self.assertTrue(labelled, p["schema"])

    def test_brief_links_come_from_the_feed(self):
        self.assertIn("https://news.example.test/metro", self.daily["mails"][0]["text"])

    # ---------- replies ----------

    def test_only_the_trusted_signed_reply_reaches_claude(self):
        self.assertEqual([p["schema"] for p in self.replies["prompts"]], ["actions"])
        self.assertIn("add milk to shopping", self.replies["prompts"][0]["prompt"])
        for text in ("pony", "caviar", "champagne"):
            self.assertNotIn(text, self.replies["prompts"][0]["prompt"])

    def test_reply_actions_happen(self):
        calls = [(c["command"], c["args"]) for c in self.home.ek_calls()]
        self.assertIn(("add-reminder", ["--list", "Shopping", "--title", "milk"]), calls)
        self.assertIn(("add-reminder", ["--list", "Bento", "--title", "Call the plumber"]), calls)
        event = next(a for c, a in calls if c == "add-event")
        self.assertEqual(event[event.index("--calendar") + 1], "Johnson Family")
        self.assertEqual(event[event.index("--start") + 1], "2026-10-15T10:00:00+01:00")
        step = (self.home.root / "projects" / "home.yaml").read_text()
        self.assertIn("status: done", step.split("- id: boiler-service")[1].split("- id:")[0])
        self.assertIn("Pat is away next Friday", (self.home.root / "data" / "household_notes.md").read_text())
        self.assertIn("Shorter news please", (self.home.root / "data" / "feedback.md").read_text())

    def test_one_confirmation_to_everyone(self):
        mails = self.replies["mails"]
        self.assertEqual(len(mails), 1)
        m = mails[0]
        self.assertEqual(m["subject"], TRUSTED_SUBJECT)
        self.assertEqual(m["in_reply_to"], self.trusted_id)
        self.assertEqual(sorted(m["to"]), ["alex@example.com", "sam@example.com"])
        for line in ("Added to Shopping: milk", "Added to Johnson Family: Dentist for Mia", "Marked done"):
            self.assertIn(line, m["text"])

    def test_every_reply_is_marked_read_and_handled_once(self):
        self.assertEqual(self.replies["imap_seen"], ["1", "2", "3", "4"])
        self.assertEqual(self.replies_again["mails"], [])
        self.assertEqual(self.replies_again["prompts"], [])

    def test_nothing_is_ever_deleted(self):
        used = {c["command"] for c in self.home.ek_calls()}
        self.assertLessEqual(used, SAFE_EK_COMMANDS, used - SAFE_EK_COMMANDS)

    # ---------- letters ----------

    def test_weekly_letter_on_the_default_sunday(self):
        mails = self.weekly["mails"]
        self.assertEqual(len(mails), 1)
        self.assertTrue(mails[0]["subject"].startswith("Bento · your week ahead"), mails[0]["subject"])
        self.assertIn("Week ending Sunday", mails[0]["text"])
        self.assertIn("Bento, Sunday letter.", mails[0]["html"])
        self.assertIn("An invented, calm week.", mails[0]["text"])
        self.assertIn("Mystery Payee Lda", mails[0]["text"], "unmapped spending should be asked about")
        # Claude's figure for an idea is labelled as Claude's, not passed off as computed.
        self.assertIn("Cancel one streaming service (Claude's rough estimate: about €13 a month).", mails[0]["text"])

    def test_monthly_look_back(self):
        mails = self.monthly["mails"]
        self.assertEqual(len(mails), 1)
        self.assertIn("looking back", mails[0]["subject"])
        self.assertIn("An invented month.", mails[0]["text"])

    # ---------- portal ----------

    def test_every_portal_page_loads(self):
        for route, r in self.portal["portal"].items():
            if route.startswith("GET") and not route.endswith("unpaired"):
                self.assertIn(r["status"], (200, 302), route)
        self.assertEqual(self.portal["portal"]["GET /"]["status"], 302)

    def test_a_browser_without_a_device_key_is_sent_to_pair(self):
        unpaired = self.portal["portal"]["GET /info unpaired"]
        self.assertEqual(unpaired["status"], 303)
        self.assertIn("/pair", unpaired["body"])

    def test_portal_changes_go_through_and_cross_site_posts_dont(self):
        p = self.portal["portal"]
        for route in ("POST /lists/add", "POST /lists/complete"):
            self.assertEqual(p[route]["status"], 303, route)
        self.assertTrue(any(k.startswith("POST /projects/") and v["status"] == 303 for k, v in p.items()))
        self.assertEqual(p["POST /lists/add cross-site"]["status"], 403)
        calls = [(c["command"], c["args"]) for c in self.home.ek_calls()]
        self.assertIn(("add-reminder", ["--list", "Shopping", "--title", "Bread"]), calls)
        self.assertNotIn(("add-reminder", ["--list", "Shopping", "--title", "x"]), calls)


class BudgetByReplyTest(unittest.TestCase):
    """The look-back suggests budget changes; a trusted, signed reply applies one to the household's
    own workbook (a copy of the example, in the test home) after a backup. Strangers and failing
    signatures change nothing."""

    APPLY, UNCLEAR, FORGED = "Re: Bento · September looking back", "Re: Bento · food", "Re: Bento · budget"

    @classmethod
    def setUpClass(cls):
        cls.home = h = Home()
        cls.book = h.root / "budget" / "budget.xlsx"
        cls.monthly = h.run("monthly")
        # The look-back saved its own suggestions; these invented ones replace them so the test
        # doesn't depend on which month the fixtures fall in.
        saved = h.root / "data" / "budget_suggestions.json"
        cls.saved = json.loads(saved.read_text())
        saved.write_text(json.dumps({"month": cls.saved["month"], "sent": str(fixtures.today() - timedelta(days=1)),
                                     "suggestions": {"Lifestyle": {"budget": 120, "trend": 13, "suggested": 25},
                                                     "Food": {"budget": 950, "trend": 1100, "suggested": 1100}}}))
        h.set_actions({
            cls.APPLY: [{"type": "budget_change", "text": "apply the Lifestyle suggestion", "category": "Lifestyle"}],
            cls.UNCLEAR: [{"type": "budget_change", "text": "Food 900", "category": "Food", "amount": "900"}],
            cls.FORGED: [{"type": "budget_change", "text": "Lifestyle 200", "category": "Lifestyle", "amount": "200"}],
        })
        h.add_reply("stranger@attacker.test", cls.FORGED, "Lifestyle 200", "attacker.test")
        h.add_reply("alex@example.com", cls.FORGED, "Lifestyle 200", None)
        h.add_reply("alex@example.com", cls.FORGED, "Lifestyle 200", "example.com.attacker.test")
        cls.before = cls.book.read_bytes()
        cls.untrusted = h.run("replies")
        cls.after_untrusted = cls.book.read_bytes()
        cls.history_after_untrusted = cls.history()
        h.add_reply("sam@example.com", cls.APPLY, "Yes, please apply the Lifestyle suggestion.", "example.com")
        h.add_reply("alex@example.com", cls.UNCLEAR, "Food 900", "example.com")
        cls.trusted = h.run("replies")

    @classmethod
    def tearDownClass(cls):
        cls.home.close()

    @classmethod
    def history(cls) -> list:
        folder = cls.home.root / "budget" / "history"
        return sorted(folder.iterdir()) if folder.exists() else []

    @staticmethod
    def amount(path, item):
        import openpyxl
        ws = openpyxl.load_workbook(path)["Expenses"]
        return next(ws[f"D{r}"].value for r in range(5, 61) if ws[f"B{r}"].value == item)

    def test_no_flow_crashed(self):
        for record in (self.monthly, self.untrusted, self.trusted):
            self.assertIsNone(record["error"], record["error"])
            self.assertEqual(record["schema_errors"], [])

    def test_the_sent_look_back_keeps_its_suggestions_for_replies(self):
        last_month = (fixtures.today().replace(day=1) - timedelta(days=1)).strftime("%Y-%m")
        self.assertIn(self.saved["month"], (last_month, fixtures.today().strftime("%Y-%m")))
        text = self.monthly["mails"][0]["text"]
        for category in self.saved["suggestions"]:
            self.assertIn(f"  {category}: budget", text)
        if self.saved["suggestions"]:
            self.assertIn("BUDGET SUGGESTIONS (REPLY TO APPLY ONE)", text)

    def test_untrusted_or_unsigned_replies_change_nothing(self):
        self.assertEqual(self.untrusted["prompts"], [])
        self.assertEqual(self.untrusted["mails"], [])
        self.assertEqual(self.after_untrusted, self.before)
        self.assertEqual(self.history_after_untrusted, [])

    def test_a_trusted_reply_changes_the_line_after_a_backup(self):
        self.assertIn("budget_change", self.trusted["prompts"][0]["prompt"])
        self.assertEqual(self.amount(self.book, "Subscriptions and outings"), 25)
        backups = self.history()
        self.assertEqual(len(backups), 1, backups)
        self.assertRegex(backups[0].name, r"^budget_\d{4}-\d{2}-\d{2}_reply\.xlsx$")
        self.assertEqual(backups[0].read_bytes(), self.before)
        self.assertEqual(self.amount(self.book, "Groceries"), 700, "the unclear reply changed Food")

    def test_confirmation_to_everyone_says_what_changed_and_where_the_backup_is(self):
        mails = {m["subject"]: m for m in self.trusted["mails"]}
        applied = mails[self.APPLY]
        self.assertEqual(sorted(applied["to"]), ["alex@example.com", "sam@example.com"])
        self.assertIn("Subscriptions and outings (Lifestyle) from €120 to €25 a month", applied["text"])
        self.assertIn(f"saved as {Path('budget', 'history', self.history()[0].name)}", applied["text"])
        unclear = mails[self.UNCLEAR]
        self.assertEqual(sorted(unclear["to"]), ["alex@example.com", "sam@example.com"])
        self.assertIn("I didn't change the budget.", unclear["text"])
        self.assertIn("Groceries, Restaurants", unclear["text"])


class WeeklyOnFridayTest(unittest.TestCase):
    """A household that moved the weekly letter to Friday morning: its wording says Friday."""

    @classmethod
    def setUpClass(cls):
        import yaml
        cls.home = h = Home()
        path = h.root / "config" / "settings.yaml"
        settings = yaml.safe_load(path.read_text())
        settings["email"]["weekly_day"], settings["email"]["weekly_hour"] = 4, 9
        path.write_text(yaml.safe_dump(settings, sort_keys=False, allow_unicode=True))
        cls.portal = h.run("portal")      # before any letter exists
        cls.weekly = h.run("weekly")

    @classmethod
    def tearDownClass(cls):
        cls.home.close()

    def test_letter_names_friday_and_never_sunday(self):
        self.assertIsNone(self.weekly["error"], self.weekly["error"])
        m = self.weekly["mails"][0]
        friday = fixtures.today() + timedelta(days=(4 - fixtures.today().weekday()) % 7)
        self.assertEqual(m["subject"], f"Bento · your week ahead, {strftime(friday + timedelta(days=1), '%-d %B')}")
        self.assertIn(f"Week ending Friday {friday.day} {friday:%B %Y}", m["text"])
        self.assertIn("Bento, Friday letter.", m["html"])
        for part in ("subject", "text", "html"):
            self.assertNotIn("Sunday", m[part], part)

    def test_claude_is_told_the_day(self):
        prompt = next(p for p in self.weekly["prompts"] if p["schema"] == "weekly")
        self.assertIn("It goes out on Friday morning.", prompt["system"])
        self.assertNotIn("Sunday", prompt["system"] + prompt["prompt"])

    def test_empty_week_page_says_friday(self):
        self.assertIn("The first weekly letter arrives Friday at 9:00.", self.portal["portal"]["GET /week"]["body"])


class MidCourseTest(unittest.TestCase):
    """Someone halfway through the course: settings written, nothing else connected yet."""

    @classmethod
    def setUpClass(cls):
        cls.home = Home(ek=False, secrets=(), projects=False, reference=False)
        cls.preview = cls.home.run("daily-dry")

    @classmethod
    def tearDownClass(cls):
        cls.home.close()

    def test_preview_builds_and_says_what_is_not_connected(self):
        self.assertIsNone(self.preview["error"], self.preview["error"])
        html = next((self.home.root / "out" / "preview").glob("brief-*.html")).read_text()
        self.assertIn("Not connected yet: calendar, money, to-dos.", html)
        self.assertNotIn("Problems this morning", html)
        self.assertNotIn("Traceback", self.preview["stderr"])

    def test_preview_sends_nothing_and_shows_no_example_projects(self):
        self.assertEqual(self.preview["mails"], [])
        html = next((self.home.root / "out" / "preview").glob("brief-*.html")).read_text()
        self.assertNotIn("Next up", html)


class NetWorthOffTest(unittest.TestCase):
    """`networth.in_emails: false` (council #218): no net worth in the brief, the letter or the
    look-back, while the snapshots go on for the portal. The Johnsons' figures are invented."""

    @classmethod
    def setUpClass(cls):
        cls.on = Home()
        cls.on_daily = cls.on.run("daily")
        cls.home = h = Home()
        h.edit_settings(lambda s: s["networth"].update(in_emails=False))
        cls.daily, cls.weekly, cls.monthly = h.run("daily"), h.run("weekly"), h.run("monthly")
        cls.portal = h.run("portal")

    @classmethod
    def tearDownClass(cls):
        cls.on.close()
        cls.home.close()

    def test_on_by_default(self):
        self.assertIn("Net worth", self.on_daily["mails"][0]["text"])

    def test_off_leaves_it_out_of_every_email(self):
        for name in ("daily", "weekly", "monthly"):
            record = getattr(self, name)
            self.assertIsNone(record["error"], f"{name}:\n{record['error']}")
            self.assertEqual(len(record["mails"]), 1, name)
            for part in ("text", "html"):
                self.assertNotIn("Net worth", record["mails"][0][part], name)
        for p in self.monthly["prompts"]:
            self.assertNotIn('"net_worth": {', p["prompt"])

    def test_the_snapshots_go_on_for_the_portal(self):
        self.assertTrue((self.home.root / "data" / "networth.jsonl").exists())
        self.assertEqual(self.portal["portal"]["GET /money"]["status"], 200)


if __name__ == "__main__":
    unittest.main()
